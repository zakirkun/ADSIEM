# server-api/app/services/report_excel.py
"""Renders a ReportDataset to a multi-sheet .xlsx for analysts who want to filter."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.services.report_data import (
    SEVERITY_ORDER, ReportDataset, label_category, label_severity, label_status,
)

_HEADER_FILL = PatternFill("solid", fgColor="1E293B")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_THIN = Side(style="thin", color="E2E8F0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_SEVERITY_FILL = {
    "critical": PatternFill("solid", fgColor="FEE2E2"),
    "high": PatternFill("solid", fgColor="FFEDD5"),
    "medium": PatternFill("solid", fgColor="FEF9C3"),
    "low": PatternFill("solid", fgColor="DBEAFE"),
    "info": PatternFill("solid", fgColor="F1F5F9"),
}


def _write_sheet(ws, headers: list[str], rows: list[list], widths: list[int] | None = None,
                 severity_column: int | None = None):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for r_idx in range(2, ws.max_row + 1):
        for cell in ws[r_idx]:
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if severity_column is not None:
            value = str(ws.cell(row=r_idx, column=severity_column).value or "").lower()
            fill = _SEVERITY_FILL.get(value)
            if fill:
                for cell in ws[r_idx]:
                    cell.fill = fill
    for idx, header in enumerate(headers, start=1):
        width = (widths[idx - 1] if widths and idx <= len(widths)
                 else max(12, min(48, len(str(header)) + 6)))
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _summary_rows(ds: ReportDataset) -> list[list]:
    rows = [
        ["Organisasi", ds.org_name or "-"],
        ["Jenis laporan", ds.title],
        ["Cakupan", ds.subject_label or "-"],
        ["Waktu dibuat", ds.generated_at.strftime("%Y-%m-%d %H:%M UTC")],
    ]
    if ds.period_from and ds.period_to:
        rows.append(["Periode", f"{ds.period_from.isoformat()} s.d. {ds.period_to.isoformat()}"])
    rows += [
        ["Total peringatan", ds.total_alerts],
        ["Rata-rata peringatan per hari", ds.avg_alerts_per_day],
        ["Total kasus", ds.total_cases],
    ]
    if ds.busiest_day:
        rows.append(["Hari tersibuk", f"{ds.busiest_day[0]} ({ds.busiest_day[1]} peringatan)"])
    for sev in SEVERITY_ORDER:
        rows.append([f"Peringatan {label_severity(sev)}", ds.severity_counts.get(sev, 0)])
    for cat, n in sorted(ds.category_counts.items(), key=lambda kv: kv[1], reverse=True):
        rows.append([f"Kategori: {label_category(cat)}", n])
    if ds.limitations:
        rows.append(["", ""])
        rows.append(["Catatan keterbatasan data", ""])
        for item in ds.limitations:
            rows.append([item, ""])
    return rows


def build_report_excel(ds: ReportDataset) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Ringkasan"
    _write_sheet(ws, ["Keterangan", "Nilai"], _summary_rows(ds), widths=[44, 60])
    ws.auto_filter.ref = None

    if ds.daily_counts:
        daily = wb.create_sheet("Volume Harian")
        _write_sheet(daily, ["Tanggal", "Jumlah peringatan"], [[d, n] for d, n in ds.daily_counts],
                     widths=[16, 20])

    if ds.timeline:
        tl = wb.create_sheet("Kronologi")
        _write_sheet(
            tl,
            ["Waktu (UTC)", "Jenis", "Kejadian", "Keterangan", "Tingkat", "IP sumber", "Host", "Aktor"],
            [[
                e.occurred_at.strftime("%Y-%m-%d %H:%M:%S"),
                e.kind,
                e.title,
                e.detail,
                label_severity(e.severity) if e.severity else "",
                e.source_ip or "",
                e.hostname or "",
                e.actor or "",
            ] for e in ds.timeline],
            widths=[19, 14, 46, 60, 12, 17, 22, 12],
            severity_column=5,
        )

    if ds.source_ips:
        ips = wb.create_sheet("Alamat IP Sumber")
        _write_sheet(
            ips,
            ["Alamat IP", "Jumlah peringatan", "Kritis", "Tinggi", "Sedang", "Rendah",
             "Host terkena", "Pertama terlihat", "Terakhir terlihat"],
            [[
                r.source_ip,
                r.alert_count,
                r.severities.get("critical", 0),
                r.severities.get("high", 0),
                r.severities.get("medium", 0),
                r.severities.get("low", 0),
                ", ".join(r.hostnames),
                r.first_seen.strftime("%Y-%m-%d %H:%M") if r.first_seen else "",
                r.last_seen.strftime("%Y-%m-%d %H:%M") if r.last_seen else "",
            ] for r in ds.source_ips],
            widths=[18, 16, 9, 9, 9, 9, 34, 18, 18],
        )

    if ds.top_rules:
        rules = wb.create_sheet("Aturan Deteksi")
        _write_sheet(rules, ["Aturan", "Jumlah"], [[t, n] for t, n in ds.top_rules], widths=[60, 12])

    if ds.cases:
        cases = wb.create_sheet("Kasus")
        _write_sheet(
            cases,
            ["ID kasus", "Judul", "Tingkat", "Status", "Tahap serangan", "Jumlah peringatan",
             "Host terdampak", "IP sumber", "Indikator", "Mulai", "Penilaian AI"],
            [[
                c.case_id,
                c.title,
                label_severity(c.severity),
                label_status(c.status),
                label_category(c.attack_family),
                c.alert_count,
                ", ".join(c.affected_hosts),
                ", ".join(c.source_ips),
                ", ".join(c.ioc_keys[:10]),
                c.created_at.strftime("%Y-%m-%d %H:%M"),
                label_status(c.verdict) if c.verdict else "",
            ] for c in ds.cases],
            widths=[38, 46, 12, 18, 30, 16, 30, 24, 40, 17, 16],
            severity_column=3,
        )
        detail = wb.create_sheet("Kronologi per Kasus")
        rows = []
        for case in ds.cases:
            for entry in case.timeline:
                rows.append([
                    case.title,
                    entry.occurred_at.strftime("%Y-%m-%d %H:%M:%S"),
                    entry.kind,
                    entry.title,
                    entry.detail,
                    label_severity(entry.severity) if entry.severity else "",
                    entry.source_ip or "",
                    entry.hostname or "",
                    entry.actor or "",
                ])
        _write_sheet(
            detail,
            ["Kasus", "Waktu (UTC)", "Jenis", "Kejadian", "Keterangan", "Tingkat", "IP sumber", "Host", "Aktor"],
            rows, widths=[38, 19, 14, 44, 60, 12, 17, 22, 12], severity_column=6,
        )

    if ds.agents:
        agents = wb.create_sheet("Agen")
        _write_sheet(
            agents,
            ["Host", "Nama", "Status", "Skor risiko", "Peringatan", "Kritis", "Tinggi",
             "Perubahan berkas", "Skor higiene", "Sistem operasi", "Port terbuka", "Versi", "Terakhir terlihat"],
            [[
                a.hostname,
                a.name,
                label_status(a.status),
                a.risk_score,
                a.alert_count,
                a.critical_count,
                a.high_count,
                a.fim_events,
                a.hygiene_score if a.hygiene_score is not None else "",
                a.os_name or "",
                a.open_ports,
                a.version or "",
                a.last_seen_at.strftime("%Y-%m-%d %H:%M") if a.last_seen_at else "",
            ] for a in ds.agents],
            widths=[24, 24, 14, 12, 12, 9, 9, 16, 12, 24, 12, 12, 18],
        )

    if ds.mitre_techniques:
        mitre = wb.create_sheet("Teknik MITRE")
        _write_sheet(mitre, ["Teknik"], [[t] for t in ds.mitre_techniques], widths=[40])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_timestamp(ds: ReportDataset) -> str:
    return ds.generated_at.strftime("%Y%m%d-%H%M")
