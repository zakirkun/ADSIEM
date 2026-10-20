# tests/server-api/test_report_data.py
from datetime import date, datetime, timezone
from types import SimpleNamespace
from uuid import uuid4

import pytest

from app.services.report_data import (
    ReportDataset, SEVERITY_ORDER, SourceIpRow, TimelineEntry, category_of,
    dataset_prompt_facts, dataset_to_dict, label_category, label_severity,
    label_status, severity_rank,
)
from app.services.report_excel import build_report_excel
from app.services.report_pdf import build_report_pdf, report_filename


def _alert(severity="high", stage="execution", ip="203.0.113.9", host="web-01", title="Suspicious shell"):
    return SimpleNamespace(
        severity=severity,
        kill_chain_stage=stage,
        source_ip=ip,
        hostname=host,
        title=title,
        status="new",
        ai_verdict=None,
        duplicate_count=0,
        mitre_techniques=["T1059"],
        created_at=datetime(2026, 10, 1, 10, 0, tzinfo=timezone.utc),
    )


def _dataset(scope="detailed", **overrides):
    base = dict(
        scope=scope,
        title="Laporan Detail Aktivitas Keamanan",
        org_name="PT Contoh",
        generated_at=datetime(2026, 10, 20, 8, 0, tzinfo=timezone.utc),
        period_from=date(2026, 10, 1),
        period_to=date(2026, 10, 7),
        total_alerts=42,
        total_cases=3,
        avg_alerts_per_day=6.0,
        busiest_day=("2026-10-03", 15),
        severity_counts={"critical": 2, "high": 10, "medium": 20, "low": 10},
        category_counts={"execution": 12, "credential_access": 30},
        status_counts={"new": 42},
        daily_counts=[("2026-10-01", 5), ("2026-10-02", 0), ("2026-10-03", 15)],
        top_rules=[("Suspicious shell", 12), ("SSH brute force", 30)],
        mitre_techniques=["T1059", "T1110"],
        source_ips=[SourceIpRow(source_ip="203.0.113.9", alert_count=30, hostnames=["web-01"])],
        timeline=[TimelineEntry(
            occurred_at=datetime(2026, 10, 3, 4, 5, tzinfo=timezone.utc),
            kind="alert", title="SSH brute force", detail="Tahap: Akses awal",
            severity="critical", source_ip="203.0.113.9", hostname="web-01",
        )],
        subject_label="2026-10-01 s.d. 2026-10-07",
    )
    base.update(overrides)
    return ReportDataset(**base)


def test_category_falls_back_to_other_when_stage_missing():
    assert category_of(_alert(stage=None)) == "other"


def test_severity_rank_orders_critical_before_info():
    assert severity_rank("critical") < severity_rank("high") < severity_rank("info")


def test_severity_rank_puts_unknown_last():
    assert severity_rank("bogus") == len(SEVERITY_ORDER)


@pytest.mark.parametrize("raw,expected", [
    ("critical", "Kritis"),
    ("high", "Tinggi"),
    ("medium", "Sedang"),
    ("low", "Rendah"),
    ("info", "Informasi"),
    (None, "Informasi"),
])
def test_label_severity_is_indonesian(raw, expected):
    assert label_severity(raw) == expected


def test_label_category_keeps_technical_term_in_parentheses():
    assert "kredensial" in label_category("credential_access")
    assert "credential access" in label_category("credential_access")


def test_label_category_unknown_passes_through():
    assert label_category("something_new") == "something_new"


@pytest.mark.parametrize("raw,expected", [
    ("new", "Baru"),
    ("open", "Terbuka"),
    ("resolved", "Selesai"),
    (None, "-"),
])
def test_label_status_is_indonesian(raw, expected):
    assert label_status(raw) == expected


def test_dataset_to_dict_serialises_datetimes_and_nested_rows():
    payload = dataset_to_dict(_dataset())
    assert payload["scope"] == "detailed"
    assert payload["generated_at"].startswith("2026-10-20T08:00")
    assert payload["timeline"][0]["occurred_at"].startswith("2026-10-03T04:05")
    assert payload["source_ips"][0]["source_ip"] == "203.0.113.9"
    assert payload["busiest_day"] == ("2026-10-03", 15)


def test_dataset_to_dict_omits_absent_period_for_case_scope():
    payload = dataset_to_dict(_dataset(scope="cases", period_from=None, period_to=None))
    assert payload["period_from"] is None
    assert payload["period_to"] is None


def test_prompt_facts_include_avg_and_top_ip():
    facts = dataset_prompt_facts(_dataset())
    assert "Rata-rata peringatan per hari: 6.0" in facts
    assert "203.0.113.9" in facts
    assert "Total peringatan: 42" in facts


def test_prompt_facts_never_emit_empty_severity_list():
    facts = dataset_prompt_facts(_dataset(severity_counts={}, category_counts={}, source_ips=[]))
    assert "tidak ada" in facts


def test_pdf_is_a_valid_document_for_every_scope():
    for scope in ("detailed", "cases", "agents"):
        blob = build_report_pdf(_dataset(scope=scope))
        assert blob.startswith(b"%PDF"), scope
        assert len(blob) > 1500, scope


def test_pdf_survives_empty_dataset():
    empty = ReportDataset(
        scope="detailed", title="Laporan", org_name="", generated_at=datetime.now(timezone.utc),
    )
    blob = build_report_pdf(empty)
    assert blob.startswith(b"%PDF")


def test_pdf_escapes_markup_in_alert_titles():
    hostile = _dataset(timeline=[TimelineEntry(
        occurred_at=datetime(2026, 10, 3, tzinfo=timezone.utc),
        kind="alert", title="<script>alert(1)</script> & co", severity="high",
    )])
    assert build_report_pdf(hostile).startswith(b"%PDF")


def test_excel_is_a_valid_workbook_with_expected_sheets():
    from io import BytesIO
    from openpyxl import load_workbook

    blob = build_report_excel(_dataset())
    assert blob.startswith(b"PK")
    wb = load_workbook(BytesIO(blob))
    assert "Ringkasan" in wb.sheetnames
    assert "Kronologi" in wb.sheetnames
    assert "Alamat IP Sumber" in wb.sheetnames
    summary = wb["Ringkasan"]
    values = [c.value for row in summary.iter_rows() for c in row if c.value]
    assert "Total peringatan" in values


def test_excel_skips_optional_sheets_when_no_data():
    from io import BytesIO
    from openpyxl import load_workbook

    minimal = ReportDataset(
        scope="detailed", title="Laporan", org_name="X",
        generated_at=datetime.now(timezone.utc),
    )
    wb = load_workbook(BytesIO(build_report_excel(minimal)))
    assert wb.sheetnames == ["Ringkasan"]


def test_report_filename_is_filesystem_safe():
    name = report_filename(_dataset(), "pdf")
    assert name.endswith(".pdf")
    assert "/" not in name and "\\" not in name and ":" not in name
    assert name.startswith("laporan-detailed-")


def test_report_filename_uses_period_for_detailed_scope():
    assert "2026-10-01-2026-10-07" in report_filename(_dataset(), "xlsx")


def test_report_filename_uses_case_count():
    from app.services.report_data import CaseBlock

    block = CaseBlock(
        case_id=str(uuid4()), title="Kampanye", severity="high", status="open",
        created_at=datetime(2026, 10, 1, tzinfo=timezone.utc), updated_at=None,
        attack_family="credential_access", affected_hosts=["web-01"], ioc_keys=[],
        verdict=None, ai_summary="", alert_count=2, timeline=[], notes=[],
        mitre=[], source_ips=[],
    )
    ds = _dataset(scope="cases", cases=[block, block, block], subject_label="3 kasus")
    assert "kasus-3" in report_filename(ds, "pdf")
